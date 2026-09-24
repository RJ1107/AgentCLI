#!/usr/bin/env python3
"""Excel (.xlsx) toolkit for the office-docs skill.

Usage:
  xlsx_tool.py info FILE                                   sheets and their used ranges
  xlsx_tool.py read FILE [--sheet NAME] [--range A1:F30] [--formulas]
                                                           cells as a Markdown table
  xlsx_tool.py from-csv OUT --csv DATA.csv [--sheet NAME]  new workbook from a CSV file
  xlsx_tool.py set FILE OUT --cell B2=100 --cell C2==B2*2  copy with cells changed

`read` shows the values Excel last calculated. A workbook written by a script and never opened
in Excel has no cached values for its formulas; --formulas shows the formulas themselves.
In `set`, a value starting with "=" is written as a formula, a number as a number, anything
else as text. Writing commands never overwrite an input or an existing file unless --force.

Needs: pip install openpyxl
"""

from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

MAX_ROWS = 200


def _openpyxl():
    try:
        import openpyxl  # noqa: PLC0415
    except ImportError:
        sys.exit("Missing Python package 'openpyxl'. Install it with: pip install openpyxl")
    return openpyxl


def _check_output(out: Path, inputs: list[Path], force: bool) -> None:
    if any(out.resolve() == item.resolve() for item in inputs):
        sys.exit("refusing to overwrite an input file; choose another output path")
    if out.exists() and not force:
        sys.exit(f"{out} already exists; pass --force to replace it")
    out.parent.mkdir(parents=True, exist_ok=True)


def _sheet(workbook, name: str | None):
    if name is None:
        return workbook.active
    if name not in workbook.sheetnames:
        sys.exit(f"no sheet named {name!r}; sheets: {', '.join(workbook.sheetnames)}")
    return workbook[name]


def _value(text: str):
    if text.startswith("="):
        return text
    for cast in (int, float):
        try:
            return cast(text)
        except ValueError:
            continue
    return text


def cmd_info(args) -> None:
    workbook = _openpyxl().load_workbook(args.file)
    for sheet in workbook.worksheets:
        print(
            f"- {sheet.title}: {sheet.dimensions} ({sheet.max_row} rows x {sheet.max_column} cols)"
        )


def cmd_read(args) -> None:
    workbook = _openpyxl().load_workbook(args.file, data_only=not args.formulas)
    sheet = _sheet(workbook, args.sheet)
    cells = sheet[args.range] if args.range else sheet.iter_rows()
    rows = [
        ["" if cell.value is None else str(cell.value).replace("\n", " ") for cell in row]
        for row in cells
    ]
    rows = [row for row in rows if any(row)]
    if not rows:
        print(f"Sheet {sheet.title!r}: no values in that range.")
        return
    shown = rows[:MAX_ROWS]
    width = max(len(row) for row in shown)
    shown = [row + [""] * (width - len(row)) for row in shown]
    print(f"Sheet {sheet.title!r}\n")
    print("| " + " | ".join(shown[0]) + " |")
    print("|" + "---|" * width)
    for row in shown[1:]:
        print("| " + " | ".join(row) + " |")
    if len(rows) > MAX_ROWS:
        print(f"\n... {len(rows) - MAX_ROWS} more rows; narrow the view with --range.")


def cmd_from_csv(args) -> None:
    openpyxl = _openpyxl()
    source, out = Path(args.csv), Path(args.out)
    _check_output(out, [source], args.force)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = args.sheet
    with source.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.reader(handle):
            sheet.append([_value(cell) for cell in row])
    workbook.save(out)
    print(f"Wrote {out}: sheet {args.sheet!r}, {sheet.max_row} rows.")


def cmd_set(args) -> None:
    openpyxl = _openpyxl()
    source, out = Path(args.file), Path(args.out)
    _check_output(out, [source], args.force)
    workbook = openpyxl.load_workbook(source)
    sheet = _sheet(workbook, args.sheet)
    for assignment in args.cell:
        ref, sep, raw = assignment.partition("=")
        if not sep:
            sys.exit(f"expected CELL=VALUE, got {assignment!r}")
        sheet[ref.strip()] = _value(raw)
    workbook.save(out)
    print(f"Wrote {out}: {len(args.cell)} cell(s) set on {sheet.title!r}.")


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawTextHelpFormatter
    )
    sub = parser.add_subparsers(dest="command", required=True)

    p = sub.add_parser("info")
    p.add_argument("file")
    p.set_defaults(handler=cmd_info)

    p = sub.add_parser("read")
    p.add_argument("file")
    p.add_argument("--sheet")
    p.add_argument("--range")
    p.add_argument("--formulas", action="store_true")
    p.set_defaults(handler=cmd_read)

    p = sub.add_parser("from-csv")
    p.add_argument("out")
    p.add_argument("--csv", required=True)
    p.add_argument("--sheet", default="Sheet1")
    p.add_argument("--force", action="store_true")
    p.set_defaults(handler=cmd_from_csv)

    p = sub.add_parser("set")
    p.add_argument("file")
    p.add_argument("out")
    p.add_argument("--sheet")
    p.add_argument("--cell", action="append", required=True)
    p.add_argument("--force", action="store_true")
    p.set_defaults(handler=cmd_set)

    args = parser.parse_args(argv)
    args.handler(args)


if __name__ == "__main__":
    main()
