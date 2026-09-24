#!/usr/bin/env python3
"""Word (.docx) toolkit for the office-docs skill.

Usage:
  docx_tool.py read FILE                         document as Markdown (headings, lists, tables)
  docx_tool.py create OUT --from NOTES.md        new document from simple Markdown
  docx_tool.py replace FILE OUT --find A --replace B   copy with text replaced

Markdown understood by `create`: "# " to "### " headings, "- " bullets, "1. " numbered items,
"| a | b |" table rows (a "|---|" separator line is skipped), **bold** inside a line, and blank
lines between paragraphs. Writing commands never overwrite an input or an existing file
unless --force is given.

Needs: pip install python-docx
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")


def _docx():
    try:
        import docx  # noqa: PLC0415
    except ImportError:
        sys.exit("Missing Python package 'python-docx'. Install it with: pip install python-docx")
    return docx


def _check_output(out: Path, inputs: list[Path], force: bool) -> None:
    if any(out.resolve() == item.resolve() for item in inputs):
        sys.exit("refusing to overwrite an input file; choose another output path")
    if out.exists() and not force:
        sys.exit(f"{out} already exists; pass --force to replace it")
    out.parent.mkdir(parents=True, exist_ok=True)


def _blocks(document):
    """Paragraphs and tables in the order they appear in the body."""

    from docx.table import Table  # noqa: PLC0415
    from docx.text.paragraph import Paragraph  # noqa: PLC0415

    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            yield Paragraph(child, document)
        elif tag == "tbl":
            yield Table(child, document)


def cmd_read(args) -> None:
    docx = _docx()
    document = docx.Document(args.file)
    lines: list[str] = []
    for block in _blocks(document):
        if block.__class__.__name__ == "Table":
            rows = [
                [cell.text.strip().replace("\n", " ") for cell in row.cells] for row in block.rows
            ]
            if rows:
                width = max(len(row) for row in rows)
                lines.append("")
                lines.append("| " + " | ".join(rows[0]) + " |")
                lines.append("|" + "---|" * width)
                lines.extend("| " + " | ".join(row) + " |" for row in rows[1:])
                lines.append("")
            continue
        text = block.text.strip()
        if not text:
            continue
        style = (block.style.name or "").lower() if block.style is not None else ""
        heading = re.match(r"heading (\d)", style)
        if heading:
            lines.append(f"{'#' * int(heading.group(1))} {text}")
        elif style == "title":
            lines.append(f"# {text}")
        elif "list number" in style:
            lines.append(f"1. {text}")
        elif "list" in style:
            lines.append(f"- {text}")
        else:
            lines.append(text)
    print("\n".join(lines) if lines else "(the document has no text)")


def _add_runs(paragraph, text: str) -> None:
    # **bold** segments become bold runs; everything else is plain.
    for index, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            paragraph.add_run(part).bold = index % 2 == 1


def cmd_create(args) -> None:
    docx = _docx()
    source = Path(args.source)
    out = Path(args.out)
    _check_output(out, [source], args.force)
    document = docx.Document()
    table_rows: list[list[str]] = []

    def flush_table() -> None:
        if not table_rows:
            return
        width = max(len(row) for row in table_rows)
        table = document.add_table(rows=len(table_rows), cols=width)
        table.style = "Table Grid"
        for r, row in enumerate(table_rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
        table_rows.clear()

    for raw in source.read_text(encoding="utf-8").splitlines():
        line = raw.rstrip()
        if line.startswith("|"):
            cells = [cell.strip() for cell in line.strip("|").split("|")]
            if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells if cell):
                table_rows.append(cells)
            continue
        flush_table()
        heading = re.match(r"(#{1,3}) (.+)", line)
        if heading:
            document.add_heading(heading.group(2), level=len(heading.group(1)))
        elif re.match(r"[-*] ", line):
            _add_runs(document.add_paragraph(style="List Bullet"), line[2:])
        elif re.match(r"\d+\. ", line):
            _add_runs(document.add_paragraph(style="List Number"), line.split(". ", 1)[1])
        elif line.strip():
            _add_runs(document.add_paragraph(), line)
    flush_table()
    document.save(out)
    print(f"Wrote {out}.")


def cmd_replace(args) -> None:
    docx = _docx()
    source, out = Path(args.file), Path(args.out)
    _check_output(out, [source], args.force)
    document = docx.Document(source)
    count = 0

    def fix(paragraph) -> None:
        nonlocal count
        if args.find not in paragraph.text:
            return
        # Replace inside single runs first so formatting survives.
        for run in paragraph.runs:
            if args.find in run.text:
                count += run.text.count(args.find)
                run.text = run.text.replace(args.find, args.replace)
        # Text split across runs (Word does this often): rewrite the paragraph into its first
        # run. The paragraph keeps its style; mixed formatting inside it is lost.
        if args.find in paragraph.text and paragraph.runs:
            count += paragraph.text.count(args.find)
            merged = paragraph.text.replace(args.find, args.replace)
            paragraph.runs[0].text = merged
            for run in paragraph.runs[1:]:
                run.text = ""

    for block in _blocks(document):
        if block.__class__.__name__ == "Table":
            for row in block.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        fix(paragraph)
        else:
            fix(block)
    document.save(out)
    print(f"Wrote {out}: {count} replacement(s) of {args.find!r}.")


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawTextHelpFormatter
    )
    sub = parser.add_subparsers(dest="command", required=True)

    p = sub.add_parser("read")
    p.add_argument("file")
    p.set_defaults(handler=cmd_read)

    p = sub.add_parser("create")
    p.add_argument("out")
    p.add_argument("--from", dest="source", required=True)
    p.add_argument("--force", action="store_true")
    p.set_defaults(handler=cmd_create)

    p = sub.add_parser("replace")
    p.add_argument("file")
    p.add_argument("out")
    p.add_argument("--find", required=True)
    p.add_argument("--replace", required=True)
    p.add_argument("--force", action="store_true")
    p.set_defaults(handler=cmd_replace)

    args = parser.parse_args(argv)
    args.handler(args)


if __name__ == "__main__":
    main()
