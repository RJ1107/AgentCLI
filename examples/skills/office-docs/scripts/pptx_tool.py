#!/usr/bin/env python3
"""PowerPoint (.pptx) toolkit for the office-docs skill.

Usage:
  pptx_tool.py read FILE                      every slide: title, text, tables, speaker notes
  pptx_tool.py create OUT --from OUTLINE.md   new deck from an outline

Outline format for `create`: each "# " line starts a slide and is its title; "- " lines are
bullets ("  - " for a second level); "> " lines become speaker notes; other lines on the first
slide become its subtitle. Writing never overwrites an existing file unless --force.

Needs: pip install python-pptx
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")


def _pptx():
    try:
        import pptx  # noqa: PLC0415
    except ImportError:
        sys.exit("Missing Python package 'python-pptx'. Install it with: pip install python-pptx")
    return pptx


def cmd_read(args) -> None:
    pptx = _pptx()
    deck = pptx.Presentation(args.file)
    for number, slide in enumerate(deck.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None else ""
        print(f"\n## Slide {number}: {title or '(no title)'}")
        for shape in slide.shapes:
            if shape == title_shape:
                continue
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs).strip()
                    if text:
                        print(f"{'  ' * paragraph.level}- {text}")
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                if rows:
                    print("| " + " | ".join(rows[0]) + " |")
                    print("|" + "---|" * len(rows[0]))
                    for row in rows[1:]:
                        print("| " + " | ".join(row) + " |")
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                print(f"> Notes: {notes}")
    print(f"\n{len(deck.slides)} slide(s).")


def _parse_outline(text: str) -> list[dict]:
    slides: list[dict] = []
    for raw in text.splitlines():
        line = raw.rstrip()
        if line.startswith("# "):
            slides.append({"title": line[2:].strip(), "bullets": [], "notes": [], "lines": []})
        elif not slides or not line.strip():
            continue
        elif line.lstrip().startswith("- "):
            level = (len(line) - len(line.lstrip())) // 2
            slides[-1]["bullets"].append((min(level, 4), line.lstrip()[2:].strip()))
        elif line.startswith("> "):
            slides[-1]["notes"].append(line[2:].strip())
        else:
            slides[-1]["lines"].append(line.strip())
    return slides


def cmd_create(args) -> None:
    pptx = _pptx()
    source, out = Path(args.source), Path(args.out)
    if out.resolve() == source.resolve():
        sys.exit("refusing to overwrite the outline; choose another output path")
    if out.exists() and not args.force:
        sys.exit(f"{out} already exists; pass --force to replace it")
    slides = _parse_outline(source.read_text(encoding="utf-8"))
    if not slides:
        sys.exit("the outline has no '# ' slide titles")
    deck = pptx.Presentation()
    title_layout, content_layout = deck.slide_layouts[0], deck.slide_layouts[1]
    for index, spec in enumerate(slides):
        is_cover = index == 0 and not spec["bullets"]
        slide = deck.slides.add_slide(title_layout if is_cover else content_layout)
        slide.shapes.title.text = spec["title"]
        body = slide.placeholders[1].text_frame
        items = [(0, line) for line in spec["lines"]] if is_cover else spec["bullets"]
        items = items or [(0, line) for line in spec["lines"]]
        for position, (level, text) in enumerate(items):
            paragraph = body.paragraphs[0] if position == 0 else body.add_paragraph()
            paragraph.text = text
            paragraph.level = level
        if spec["notes"]:
            slide.notes_slide.notes_text_frame.text = "\n".join(spec["notes"])
    out.parent.mkdir(parents=True, exist_ok=True)
    deck.save(out)
    print(f"Wrote {out}: {len(slides)} slide(s).")


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawTextHelpFormatter
    )
    sub = parser.add_subparsers(dest="command", required=True)

    p = sub.add_parser("read")
    p.add_argument("file")
    p.set_defaults(handler=cmd_read)

    p = sub.add_parser("create")
    p.add_argument("out")
    p.add_argument("--from", dest="source", required=True)
    p.add_argument("--force", action="store_true")
    p.set_defaults(handler=cmd_create)

    args = parser.parse_args(argv)
    args.handler(args)


if __name__ == "__main__":
    main()
